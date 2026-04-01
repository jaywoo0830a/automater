"""
gui/excel_io.py
---------------
섹션별 Excel 가져오기 / 내보내기.

각 탭(계정·키워드·풀·제목)에서 독립적으로 호출한다.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# ── column definitions ─────────────────────────────────────────────

ACCOUNT_COLUMNS = [
    ("username", "아이디"),
    ("password", "비밀번호"),
    ("blog_id", "블로그 ID"),
    ("weight", "가중치"),
    ("min_posts", "최소"),
    ("max_posts", "최대"),
    ("proxy", "프록시"),
]
ACCOUNT_DEFAULTS = ["", "", "", "1", "0", "0", ""]

KV_COLUMNS = [
    ("value", "값"),
]

TITLE_COLUMNS = [
    ("title", "제목 템플릿"),
]

# ── helpers ────────────────────────────────────────────────────────


def _set_column_widths(ws, columns: list[tuple[str, str]]) -> None:
    for col_idx in range(1, len(columns) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 22


def _read_rows(ws) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        vals = [str(c).strip() if c is not None else "" for c in row]
        if any(vals):
            rows.append(vals)
    return rows


def _single_sheet_wb(title: str, columns: list[tuple[str, str]]) -> tuple[Workbook, Worksheet]:
    wb = Workbook()
    ws: Worksheet = wb.active  # type: ignore[assignment]
    ws.title = title
    _set_column_widths(ws, columns)
    return wb, ws


# ── accounts ───────────────────────────────────────────────────────

def export_accounts(path: str | Path, accounts: list[dict]) -> None:
    wb, ws = _single_sheet_wb("accounts", ACCOUNT_COLUMNS)
    for acc in accounts:
        ws.append([acc.get(k, d) for (k, _), d in zip(ACCOUNT_COLUMNS, ACCOUNT_DEFAULTS)])
    wb.save(str(path))


def import_accounts(path: str | Path) -> list[dict]:
    wb = load_workbook(str(path), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    keys = [k for k, _ in ACCOUNT_COLUMNS]
    result: list[dict] = []
    for row in _read_rows(ws):
        acc: dict = {}
        for i, key in enumerate(keys):
            val = row[i] if i < len(row) else ACCOUNT_DEFAULTS[i]
            if not val:
                val = ACCOUNT_DEFAULTS[i]
            if key in ("weight", "min_posts", "max_posts"):
                try:
                    ival = int(float(val))
                except (ValueError, TypeError):
                    ival = int(ACCOUNT_DEFAULTS[i])
                if key == "weight" and ival == 1:
                    continue
                if key in ("min_posts", "max_posts") and ival == 0:
                    continue
                acc[key] = ival
            else:
                if val:
                    acc[key] = val
        if acc.get("username"):
            result.append(acc)
    wb.close()
    return result


def template_accounts(path: str | Path) -> None:
    wb, ws = _single_sheet_wb("accounts", ACCOUNT_COLUMNS)
    ws.append(["example_id", "password123", "myblog", 1, 0, 0, ""])
    wb.save(str(path))


# ── key-value (keywords / pools) ──────────────────────────────────

def export_kv_single(path: str | Path, slug: str, values: list[str]) -> None:
    """카테고리 하나를 세로 형태 엑셀로 내보내기. 파일명 = 카테고리명."""
    wb, ws = _single_sheet_wb(slug, KV_COLUMNS)
    for val in values:
        ws.append([val])
    wb.save(str(path))


def import_kv_single(path: str | Path) -> list[str]:
    """세로 형태 엑셀에서 값 목록을 읽기. 카테고리명은 호출자가 결정."""
    wb = load_workbook(str(path), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    result: list[str] = []
    for row in _read_rows(ws):
        val = row[0] if row else ""
        if val:
            result.append(val)
    wb.close()
    return result


def template_kv_single(path: str | Path, slug: str = "category",
                        examples: list[str] | None = None) -> None:
    """카테고리 하나의 템플릿 엑셀 생성."""
    wb, ws = _single_sheet_wb(slug, KV_COLUMNS)
    for val in (examples or ["값1", "값2", "값3"]):
        ws.append([val])
    wb.save(str(path))


# ── titles ─────────────────────────────────────────────────────────

def export_titles(path: str | Path, titles: list[str]) -> None:
    wb, ws = _single_sheet_wb("titles", TITLE_COLUMNS)
    for t in titles:
        ws.append([t])
    wb.save(str(path))


def import_titles(path: str | Path) -> list[str]:
    wb = load_workbook(str(path), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    result: list[str] = []
    for row in _read_rows(ws):
        val = row[0] if row else ""
        if val:
            result.append(val)
    wb.close()
    return result


def template_titles(path: str | Path) -> None:
    wb, ws = _single_sheet_wb("titles", TITLE_COLUMNS)
    ws.append(["{keyword:region} {keyword:subject} 과외 추천"])
    wb.save(str(path))


# ── maps (key → value) ────────────────────────────────────────────

MAP_COLUMNS = [
    ("key", "키"),
    ("value", "값"),
]


def export_map(path: str | Path, data: dict[str, str | list[str]]) -> None:
    """맵 데이터를 엑셀로 내보내기. 1:N 값은 여러 행으로 확장."""
    wb, ws = _single_sheet_wb("map", MAP_COLUMNS)
    for k, v in data.items():
        if isinstance(v, list):
            for i, item in enumerate(v):
                ws.append([k if i == 0 else "", item])
        else:
            ws.append([k, v])
    wb.save(str(path))


def import_map(path: str | Path) -> dict[str, str | list[str]]:
    """맵 엑셀 읽기. 1:N 지원 -- 키가 비어있으면 이전 키에 값 추가."""
    wb = load_workbook(str(path), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    result: dict[str, list[str]] = {}
    last_key = ""
    for row in _read_rows(ws):
        key = row[0] if len(row) > 0 else ""
        val = row[1] if len(row) > 1 else ""
        if not val:
            continue
        if key:
            last_key = key
        if not last_key:
            continue
        if last_key not in result:
            result[last_key] = []
        result[last_key].append(val)
    wb.close()
    # 단일 값은 문자열로, 복수 값은 리스트로
    final: dict[str, str | list[str]] = {}
    for k, vals in result.items():
        final[k] = vals[0] if len(vals) == 1 else vals
    return final


def template_map(path: str | Path) -> None:
    wb, ws = _single_sheet_wb("map", MAP_COLUMNS)
    ws.append(["강남", "gangnam.jpg"])
    ws.append(["서초", "seocho.jpg"])
    ws.append(["다산동", "다산점1.jpg"])
    ws.append(["", "다산점2.jpg"])
    ws.append(["", "다산점3.jpg"])
    ws.append(["_default", "default.jpg"])
    wb.save(str(path))
