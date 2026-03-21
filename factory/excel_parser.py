"""
factory/excel_parser.py
-------------------------
Excel-to-dict parsing for bulk import.

Two layers:
    parse_rows(headers, rows)  — pure function, testable without I/O
    parse_excel(path)          — thin wrapper using openpyxl

Header resolution:
    resolve_headers(data, header_map, category_names)
        — maps display names ('지역') to slugs ('region')
"""

from __future__ import annotations

import re


# ---------------------------------------------------------------------------
# Slug detection — is this already a valid ASCII slug?
# ---------------------------------------------------------------------------

_SLUG_RE = re.compile(r"^[a-z][a-z0-9_]*$")


def _is_slug(value: str) -> bool:
    """Return True if value looks like a valid template slug."""
    return bool(_SLUG_RE.match(value))


# ---------------------------------------------------------------------------
# parse_rows — pure function, no I/O
# ---------------------------------------------------------------------------

def parse_rows(
    headers: list[str],
    rows: list[list],
) -> dict[str, list[str]]:
    """
    Convert columnar data into {header: [values]} dict.

    Rules:
        - Headers are stripped of whitespace
        - Empty/None cells are skipped
        - Duplicate values within a column are removed (order preserved)
        - Numeric values are converted to strings (int 1 → "1", float 1.0 → "1")
        - Columns where all cells are empty are excluded
        - Rows shorter than headers are padded with None

    Args:
        headers: Column header strings.
        rows:    List of rows, each a list of cell values.

    Returns:
        {header: [unique_values]} — only columns with at least one value.

    Raises:
        ValueError: If headers list is empty or contains blank headers.
    """
    if not headers:
        raise ValueError("headers must not be empty")

    clean_headers = [h.strip() if isinstance(h, str) else str(h) for h in headers]
    for h in clean_headers:
        if not h:
            raise ValueError("all headers must be non-empty strings")

    columns: dict[str, list[str]] = {h: [] for h in clean_headers}
    seen: dict[str, set[str]] = {h: set() for h in clean_headers}

    for row in rows:
        for i, header in enumerate(clean_headers):
            cell = row[i] if i < len(row) else None
            value = _normalize_cell(cell)
            if value is None:
                continue
            if value in seen[header]:
                continue
            columns[header].append(value)
            seen[header].add(value)

    return {h: vals for h, vals in columns.items() if vals}


def _normalize_cell(cell) -> str | None:
    """Convert a cell value to a clean string, or None if empty."""
    if cell is None:
        return None
    if isinstance(cell, float):
        cell = int(cell) if cell == int(cell) else cell
    value = str(cell).strip()
    return value if value else None


# ---------------------------------------------------------------------------
# parse_excel — openpyxl wrapper
# ---------------------------------------------------------------------------

def parse_excel(path: str, sheet_name: str | None = None) -> dict[str, list[str]]:
    """
    Parse an .xlsx file into {header: [values]} dict.

    Args:
        path:       Path to .xlsx file.
        sheet_name: Sheet to read. None = active sheet.

    Returns:
        Parsed columnar data via parse_rows.
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active

    row_iter = ws.iter_rows(values_only=True)
    header_row = next(row_iter, None)
    if not header_row:
        raise ValueError("Excel file has no header row")

    headers = [str(h).strip() if h else "" for h in header_row]
    rows = [list(r) for r in row_iter]

    wb.close()
    return parse_rows(headers, rows)


# ---------------------------------------------------------------------------
# resolve_headers — map display names to slugs
# ---------------------------------------------------------------------------

def resolve_headers(
    data: dict[str, list[str]],
    header_map: dict[str, str] | None = None,
    category_names: dict[str, str] | None = None,
) -> dict[str, list[str]]:
    """
    Map display-name headers to slug keys.

    Resolution order (first match wins):
        1. Explicit header_map: {"지역": "region"}
        2. Category name lookup: KeywordCategory(name="지역").slug → "region"
        3. Already a slug: "region" passes through unchanged

    Args:
        data:           {display_name: [values]} from parse_rows.
        header_map:     Optional explicit mapping.
        category_names: {name: slug} lookup from KeywordCategory table.

    Returns:
        {slug: [values]} — same values, keys replaced with slugs.

    Raises:
        ValueError: If any header cannot be resolved to a slug.
    """
    if not data:
        raise ValueError("data must not be empty")

    header_map = header_map or {}
    category_names = category_names or {}
    result: dict[str, list[str]] = {}
    unresolved: list[str] = []

    for header, values in data.items():
        slug = _resolve_one(header, header_map, category_names)
        if slug is None:
            unresolved.append(header)
        else:
            result[slug] = values

    if unresolved:
        raise ValueError(
            f"cannot resolve headers to slugs: {unresolved}. "
            f"Provide a header_map or register KeywordCategory "
            f"with matching name."
        )

    return result


def _resolve_one(
    header: str,
    header_map: dict[str, str],
    category_names: dict[str, str],
) -> str | None:
    """Resolve a single header to a slug. Returns None if unresolvable."""
    if header in header_map:
        return header_map[header]
    if header in category_names:
        return category_names[header]
    if _is_slug(header):
        return header
    return None
