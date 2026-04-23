"""
cli/report.py
--------------
캠페인 완료 후 발행 기록을 파일로 내보낸다.

랭크 추적 콜렉터가 소비할 수 있도록 각 발행 건의 핵심 정보만 기록한다:
    blog_id | keyword | title | published_at

포맷은 파일 확장자로 결정한다:
    .xlsx        → Excel (openpyxl)
    .yaml/.yml   → YAML
"""

from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

log = logging.getLogger(__name__)


# 랭크 콜렉터가 필요로 하는 필드만.
_FIELDS: tuple[str, ...] = ("blog_id", "keyword", "title", "published_at")


def _iter_rows(result: Any, include_failed: bool) -> Iterable[dict[str, Any]]:
    """ExecutionResult 에서 리포트 row 를 만든다."""
    records = list(result.succeeded_combos)
    if include_failed:
        records.extend(result.failed_combos)

    for rec in records:
        published = getattr(rec, "published_at", None)
        yield {
            "blog_id":      getattr(rec, "blog_id", "") or "",
            "keyword":      getattr(rec, "keyword", "") or "",
            "title":        getattr(rec, "title", "") or "",
            "published_at": published.isoformat() if isinstance(published, datetime) else "",
        }


def write_rows(rows: list[dict[str, Any]], path: str | Path) -> Path:
    """이미 평탄화된 row dict 리스트를 path 에 저장한다.

    GUI 등 ExecutionResult 없이 기존 리포트 파일을 다른 포맷으로
    재저장하려는 경우에 사용한다. 포맷은 확장자로 결정한다.
    """
    out = Path(path).expanduser().resolve()
    out.parent.mkdir(parents=True, exist_ok=True)
    ext = out.suffix.lower()

    if ext == ".xlsx":
        _write_xlsx(out, rows)
    elif ext in (".yaml", ".yml"):
        _write_yaml(out, rows)
    else:
        raise ValueError(
            f"지원하지 않는 리포트 확장자: {ext or '(없음)'} — .xlsx / .yaml / .yml 중 하나를 사용하세요."
        )

    log.info("[report] %d건 저장 완료 → %s", len(rows), out)
    return out


def write_report(
    result: Any,
    path: str | Path,
    *,
    include_failed: bool = False,
) -> Path:
    """ExecutionResult 로부터 row 를 뽑아 path 에 저장한다.

    Args:
        result:         ExecutionResult (succeeded_combos / failed_combos 를 가진 객체).
        path:           저장 경로. .xlsx / .yaml / .yml 지원.
        include_failed: True 이면 실패한 조합도 포함.

    Returns:
        저장된 파일의 절대 경로.
    """
    rows = list(_iter_rows(result, include_failed=include_failed))
    return write_rows(rows, path)


# ---------------------------------------------------------------------------
# Writers
# ---------------------------------------------------------------------------

def _write_xlsx(out: Path, rows: list[dict[str, Any]]) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "posts"
    ws.append(list(_FIELDS))
    for row in rows:
        ws.append([row[f] for f in _FIELDS])

    # 간단한 열 폭 가이드.
    widths = {"blog_id": 20, "keyword": 28, "title": 40, "published_at": 22}
    for idx, field in enumerate(_FIELDS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = widths[field]

    wb.save(out)


def _write_yaml(out: Path, rows: list[dict[str, Any]]) -> None:
    import yaml

    with open(out, "w", encoding="utf-8") as f:
        yaml.safe_dump(
            rows,
            f,
            allow_unicode=True,
            default_flow_style=False,
            sort_keys=False,
        )
